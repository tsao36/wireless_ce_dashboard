"""Serve the latest weighted team loading snapshot."""

from __future__ import annotations

import argparse
import base64
import csv
import json
import os
import re
import subprocess
import threading
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, time, timedelta
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

import psycopg2
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from psycopg2 import sql

from APIs import Sherlock
import issue_category_model as category_model

ROOT = Path(__file__).resolve().parent
SNAPSHOT_PATH = Path(os.environ.get("OFFLOAD_LOADING_SNAPSHOT_FILE", "team_loading_latest.json"))
if not SNAPSHOT_PATH.is_absolute():
    SNAPSHOT_PATH = ROOT / SNAPSHOT_PATH
ACTIVITY_WORKBOOK_PATH = Path(os.environ.get("CFE_ACTIVITY_WORKBOOK_FILE", str(ROOT / "CFE work overview.xlsx")))
if not ACTIVITY_WORKBOOK_PATH.is_absolute():
    ACTIVITY_WORKBOOK_PATH = ROOT / ACTIVITY_WORKBOOK_PATH
ACTIVITY_WORKBOOK_USER = os.environ.get("CFE_ACTIVITY_WORKBOOK_USER", "").strip()
ACTIVITY_WORKBOOK_PASSWORD = os.environ.get("CFE_ACTIVITY_WORKBOOK_PASSWORD", "")
ACTIVITY_WORKBOOK_URL = os.environ.get("CFE_ACTIVITY_WORKBOOK_URL", "").strip()
ACTIVITY_CACHE: dict = {"key": None, "payload": None, "checked_at": 0.0, "refreshing": False}
ACTIVITY_LOCK = threading.Lock()
# Checking the workbook on the network share costs ~1s, so only look for changes this often.
ACTIVITY_CHECK_SECONDS = 60
ACTIVITY_WORKBOOK_CACHE_PATH = ROOT / "cache" / "CFE work overview.xlsx"
ACTIVITY_WORKBOOK_REFRESH_HOURS = float(os.environ.get("CFE_ACTIVITY_WORKBOOK_REFRESH_HOURS", "6"))
# Only local copies can be served by the static handler; resolving a UNC path is a slow network call.
BLOCKED_STATIC_PATHS = {
    path.resolve()
    for path in (ACTIVITY_WORKBOOK_PATH, ACTIVITY_WORKBOOK_CACHE_PATH)
    if not str(path).startswith("\\\\")
}
CURRENT_YEAR = datetime.now().year
HISTORY_PATH = Path(os.environ.get("OFFLOAD_LOADING_HISTORY_DIR", str(ROOT / "loading_history")))
if not HISTORY_PATH.is_absolute():
    HISTORY_PATH = ROOT / HISTORY_PATH
WEIGHT_MAP_PATH = ROOT / "issue_category_weights.json"
BATCH_PATH = ROOT / "run_offload_loading_summary_daily.bat"
RESTART_SCRIPT_PATH = ROOT / "restart_dashboard.ps1"
RESTART_COOLDOWN = timedelta(minutes=2)
RESTART_REQUESTED_AT: datetime | None = None
RUN_LOCK = threading.Lock()
RUN_PROCESS: subprocess.Popen[bytes] | None = None
RUN_STARTED_AT: str | None = None
RUN_FINISHED_AT: str | None = None
RUN_RETURN_CODE: int | None = None
FIELD_ISSUE_TABLE = "ips_jira_bugs"
FIELD_ISSUE_KEYWORD = "field issue"
FIELD_ISSUE_DETAIL_COLUMNS = (
    "ips_case_number",
    "ips_title",
    "ips_status",
    "ips_url",
    "jira_id",
    "jira_summary",
    "jira_status",
    "jira_url",
    "customer",
)
CATEGORY_TUNING_DIR = Path(
    os.environ.get(
        "CATEGORY_TUNING_DIR",
        str(ROOT.parents[1] / "jira_ips_tat_HP" / "jira_tat" / "jira_customer_tat" / "jobs" / "01_issue_category_tuning"),
    )
)
_TUNING_MODEL_PATH = CATEGORY_TUNING_DIR / "models" / "issue_category_model.joblib"
CATEGORY_MODEL_PATH = _TUNING_MODEL_PATH if _TUNING_MODEL_PATH.is_file() else ROOT / "models" / "issue_category_model.joblib"
CATEGORY_METRICS_PATH = CATEGORY_TUNING_DIR / "models" / "issue_category_model_metrics.json"
CATEGORY_CONFIG_PATH = CATEGORY_TUNING_DIR / "bug_category_config.json"
# Saved into CFE_input so the weekly tuning/retrain jobs pick team reviews up as training rows.
CATEGORY_REVIEW_PATH = Path(
    os.environ.get(
        "CATEGORY_REVIEW_FILE",
        str(
            CATEGORY_TUNING_DIR / "CFE_input" / "dashboard_cfe_reviews.csv"
            if (CATEGORY_TUNING_DIR / "CFE_input").is_dir()
            else ROOT / "category_reviews" / "dashboard_cfe_reviews.csv"
        ),
    )
)
CATEGORY_REVIEW_COLUMNS = (
    "ips_title",
    "predicted_category",
    "technology",
    "human_category",
    "ips_case_number",
    "verdict",
    "llm_category",
    "reviewer",
    "reviewed_at",
)
CATEGORY_LOCK = threading.Lock()
CATEGORY_CACHE: dict = {"model_mtime": None, "bundle": None, "predictions": {}, "queue": {}, "reviewers": set()}
CATEGORY_QUEUE_STATE: dict = {"loaded_at": 0.0, "refreshing": False}
CATEGORY_QUEUE_LOCK = threading.Lock()
# The issue query takes ~3s; serve the last result and refresh it in the background after this many seconds.
CATEGORY_QUEUE_MAX_AGE = 300
LLM_CACHE_PATH = ROOT / "cache" / "category_llm_predictions.json"
LLM_LOCK = threading.Lock()
LLM_STATE: dict = {"running": False, "done": 0, "total": 0, "failed": 0, "predictions": None, "error": "", "failed_at": None}
LLM_RETRY_AFTER = timedelta(minutes=30)
LLM_MAX_CONSECUTIVE_FAILURES = 8


def _workbook_value(value: object) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return value


def _is_current_year(value: object) -> bool:
    return bool(re.search(rf"(?<!\d){CURRENT_YEAR}(?!\d)", str(value or "")))


def _activity_category(sheet_name: str, headers: tuple[object, ...]) -> str:
    name = sheet_name.casefold()
    if any(term in name for term in ("budget", "shopping", "expense")):
        return "Budget & purchasing"
    if any(term in name for term in ("training", "skill", "evaluation")):
        return "Learning & skills"
    if (
        any(term in name for term in ("meeting", "rotation", "schedule", "presentation"))
        or (headers and str(headers[0]).casefold() == "ww/time")
    ):
        return "Schedules & events"
    if any(term in name for term in ("task", "project", "feature", "rnr")):
        return "Projects & tasks"
    return "Team reference"


def _ensure_activity_workbook() -> Path:
    if ACTIVITY_WORKBOOK_PATH.is_file():
        return ACTIVITY_WORKBOOK_PATH
    if str(ACTIVITY_WORKBOOK_PATH).startswith("\\\\"):
        _connect_network_workbook()
    if ACTIVITY_WORKBOOK_PATH.is_file():
        return ACTIVITY_WORKBOOK_PATH
    if not ACTIVITY_WORKBOOK_URL:
        raise FileNotFoundError(ACTIVITY_WORKBOOK_PATH.name)

    cache_path = ACTIVITY_WORKBOOK_CACHE_PATH
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_is_fresh = cache_path.is_file() and (
        datetime.now().timestamp() - cache_path.stat().st_mtime < ACTIVITY_WORKBOOK_REFRESH_HOURS * 3600
    )
    if cache_is_fresh:
        return cache_path

    request = urllib.request.Request(
        ACTIVITY_WORKBOOK_URL,
        headers={"User-Agent": "Wireless-CFE-Dashboard/1.0"},
    )
    temp_path = cache_path.with_suffix(".tmp")
    try:
        try:
            with urllib.request.urlopen(request, timeout=45) as response:
                content = response.read(50 * 1024 * 1024 + 1)
        except urllib.error.HTTPError:
            content = b""
        if not content.startswith(b"PK"):
            content = _download_sharepoint_via_graph(ACTIVITY_WORKBOOK_URL)
        if len(content) > 50 * 1024 * 1024:
            raise ValueError("Workbook exceeds the 50 MB download limit.")
        temp_path.write_bytes(content)
        os.replace(temp_path, cache_path)
        return cache_path
    except (OSError, urllib.error.URLError, ValueError) as exc:
        temp_path.unlink(missing_ok=True)
        if cache_path.is_file():
            return cache_path
        raise FileNotFoundError(f"Unable to download workbook from SharePoint: {exc}") from exc


def _connect_network_workbook() -> None:
    if not ACTIVITY_WORKBOOK_USER or not ACTIVITY_WORKBOOK_PASSWORD:
        return
    match = re.match(r"^(\\\\[^\\]+\\[^\\]+)", str(ACTIVITY_WORKBOOK_PATH))
    if not match:
        raise ValueError("CFE_ACTIVITY_WORKBOOK_FILE is not a valid UNC path.")
    result = subprocess.run(
        [
            "net",
            "use",
            match.group(1),
            f"/user:{ACTIVITY_WORKBOOK_USER}",
            ACTIVITY_WORKBOOK_PASSWORD,
            "/persistent:no",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0 and not ACTIVITY_WORKBOOK_PATH.is_file():
        raise ValueError(f"Unable to connect to workbook share ({result.returncode}).")


def _download_sharepoint_via_graph(share_url: str) -> bytes:
    tenant = os.environ.get("AZURE_TENANT_ID", "").strip()
    client_id = os.environ.get("AZURE_CLIENT_ID", "").strip()
    client_secret = os.environ.get("GRAPH_CLIENT_SECRET", "").strip()
    if not all((tenant, client_id, client_secret)):
        raise ValueError("SharePoint requires AZURE_TENANT_ID, AZURE_CLIENT_ID, and GRAPH_CLIENT_SECRET.")

    token_request = urllib.request.Request(
        f"https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token",
        data=(
            f"client_id={quote(client_id)}&"
            f"client_secret={quote(client_secret)}&"
            "scope=https%3A%2F%2Fgraph.microsoft.com%2F.default&"
            "grant_type=client_credentials"
        ).encode("utf-8"),
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        method="POST",
    )
    with urllib.request.urlopen(token_request, timeout=30) as response:
        token_payload = json.loads(response.read().decode("utf-8"))
    access_token = token_payload.get("access_token")
    if not access_token:
        raise ValueError("Microsoft Graph did not return an access token.")

    share_id = "u!" + base64.urlsafe_b64encode(share_url.encode("utf-8")).decode("ascii").rstrip("=")
    content_request = urllib.request.Request(
        f"https://graph.microsoft.com/v1.0/shares/{share_id}/driveItem/content",
        headers={"Authorization": f"Bearer {access_token}"},
    )
    with urllib.request.urlopen(content_request, timeout=60) as response:
        content = response.read(50 * 1024 * 1024 + 1)
    if not content.startswith(b"PK"):
        raise ValueError("Microsoft Graph returned a non-Excel response.")
    return content


def _refresh_activity_workbook() -> None:
    try:
        _load_activity_workbook(background=True)
    except Exception as exc:
        print(f"[dashboard] Activity workbook refresh failed ({type(exc).__name__}); keeping previous data.")
    finally:
        with ACTIVITY_LOCK:
            ACTIVITY_CACHE["checked_at"] = datetime.now().timestamp()
            ACTIVITY_CACHE["refreshing"] = False


def _load_activity_workbook(force: bool = False, background: bool = False) -> dict:
    now = datetime.now().timestamp()
    with ACTIVITY_LOCK:
        if not force and not background and ACTIVITY_CACHE["payload"]:
            if now - ACTIVITY_CACHE["checked_at"] >= ACTIVITY_CHECK_SECONDS and not ACTIVITY_CACHE["refreshing"]:
                ACTIVITY_CACHE["refreshing"] = True
                threading.Thread(target=_refresh_activity_workbook, daemon=True).start()
            return ACTIVITY_CACHE["payload"]
    workbook_path = _ensure_activity_workbook()
    stat = workbook_path.stat()
    # Re-parse only when the workbook file actually changes; parsing takes seconds.
    cache_key = (str(workbook_path), stat.st_mtime, stat.st_size)
    with ACTIVITY_LOCK:
        if not force and ACTIVITY_CACHE["key"] == cache_key:
            ACTIVITY_CACHE["checked_at"] = now
            return ACTIVITY_CACHE["payload"]

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            sheet_years = re.findall(r"(?<!\d)(20\d{2})(?!\d)", worksheet.title)
            if sheet_years and CURRENT_YEAR not in {int(year) for year in sheet_years}:
                continue

            source_rows = list(worksheet.iter_rows(values_only=True))
            if not source_rows:
                continue
            header_row = tuple(source_rows[0])
            year_column = next(
                (index for index, value in enumerate(header_row) if str(value or "").strip().casefold() == "year"),
                None,
            )
            selected_columns = list(range(len(header_row)))
            if worksheet.title.casefold() == "cirm host rotation":
                selected_columns = [
                    index for index, value in enumerate(header_row)
                    if index == 0 or _is_current_year(value)
                ]

            rows = []
            for row_number, source_row in enumerate(source_rows, start=1):
                if row_number > 1 and year_column is not None:
                    if year_column >= len(source_row) or not _is_current_year(source_row[year_column]):
                        continue
                values = [
                    _workbook_value(source_row[index]) if index < len(source_row) else None
                    for index in selected_columns
                ]
                if any(value is not None for value in values):
                    rows.append({"row": row_number, "values": values})

            active_columns = [
                position for position in range(len(selected_columns))
                if any(row["values"][position] is not None for row in rows)
            ]
            selected_columns = [selected_columns[position] for position in active_columns]
            rows = [
                {"row": row["row"], "values": [row["values"][position] for position in active_columns]}
                for row in rows
            ]
            headers = tuple(header_row[index] if index < len(header_row) else None for index in selected_columns)
            sheets.append(
                {
                    "name": worksheet.title,
                    "category": _activity_category(worksheet.title, headers),
                    "columns": [
                        {
                            "letter": get_column_letter(index + 1),
                            "label": str(value).strip() if value is not None else get_column_letter(index + 1),
                        }
                        for index, value in zip(selected_columns, headers)
                    ],
                    "rows": rows,
                    "row_count": max(0, len(rows) - 1),
                }
            )
    finally:
        workbook.close()

    payload = {
        "available": True,
        "year": CURRENT_YEAR,
        "source": workbook_path.name,
        "source_updated_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        "loaded_at": datetime.now().isoformat(timespec="seconds"),
        "sheets": sheets,
    }
    with ACTIVITY_LOCK:
        ACTIVITY_CACHE.update(key=cache_key, payload=payload, checked_at=now)
    return payload


def _field_issue_text(value: object) -> str:
    text = "" if value is None else str(value).strip()
    return "" if text.upper() == "NA" else text


def _field_issue_active_expr(available: set[str]) -> str:
    def status(column: str) -> str:
        if column not in available:
            return "NULL"
        return f"NULLIF(NULLIF(UPPER(TRIM({column}::text)), ''), 'NA')"

    bug, ips, jira = status("bug_status"), status("ips_status"), status("jira_status")
    # bug_status wins; otherwise closed if any present IPS/Jira status is closed, and no status at all is not active.
    return (
        f"(CASE WHEN {bug} IS NOT NULL THEN {bug} <> 'CLOSED' "
        f"WHEN {ips} IS NULL AND {jira} IS NULL THEN FALSE "
        f"ELSE COALESCE({ips}, '') <> 'CLOSED' AND COALESCE({jira}, '') NOT IN ('CLOSED', 'VERIFY', 'IMPLEMENTED') END)"
    )


def _load_field_issues() -> dict:
    db = Sherlock.PostgresCustomerEngineeringDb
    connection = psycopg2.connect(
        dbname=db.database, user=db.user, password=db.password, host=db.host, port=db.port, connect_timeout=10
    )
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT column_name FROM information_schema.columns "
                "WHERE table_name = %s AND table_schema = ANY(current_schemas(false))",
                (FIELD_ISSUE_TABLE,),
            )
            available = {row[0] for row in cursor.fetchall()}
            columns = ["engineer", *(name for name in FIELD_ISSUE_DETAIL_COLUMNS if name in available)]
            pattern = f"%{FIELD_ISSUE_KEYWORD}%"
            # bug_created_date mirrors vw_issues, which maps it to ips_created_date.
            cursor.execute(
                sql.SQL(
                    "SELECT {columns}, {active} AS is_active, ips_created_date AS bug_created_date FROM {table} "
                    "WHERE ips_title::text ILIKE %s OR jira_summary::text ILIKE %s"
                ).format(
                    columns=sql.SQL(", ").join(map(sql.Identifier, columns)),
                    active=sql.SQL(_field_issue_active_expr(available)),
                    table=sql.Identifier(FIELD_ISSUE_TABLE),
                ),
                (pattern, pattern),
            )
            rows = cursor.fetchall()
    finally:
        connection.close()

    issues = []
    for *values, is_active, created in rows:
        issue = {name: _field_issue_text(value) for name, value in zip(columns, values)}
        issue["engineer"] = issue["engineer"] or "Unassigned"
        issue["active"] = bool(is_active)
        issue["bug_created_date"] = created.date().isoformat() if created else ""
        issue["in_year"] = bool(created) and created.year == CURRENT_YEAR
        issues.append(issue)
    return {
        "keyword": FIELD_ISSUE_KEYWORD,
        "year": CURRENT_YEAR,
        "loaded_at": datetime.now().isoformat(timespec="seconds"),
        "issues": issues,
    }


def _read_json(path: Path) -> dict:
    try:
        with path.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def _file_date(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime).date().isoformat() if path.exists() else ""


def _category_labels() -> list[str]:
    labels = _read_json(CATEGORY_CONFIG_PATH).get("default_categories") or []
    if not labels and CATEGORY_CACHE["bundle"]:
        labels = CATEGORY_CACHE["bundle"].get("labels") or []
    return sorted({str(label).strip() for label in labels if str(label).strip()}, key=str.casefold)


def _category_model_status() -> dict:
    metrics = _read_json(CATEGORY_METRICS_PATH)
    report = metrics.get("classification_report") or {}
    per_category = [
        {
            "category": name,
            "precision": values.get("precision"),
            "recall": values.get("recall"),
            "f1": values.get("f1-score"),
            "support": values.get("support"),
        }
        for name, values in report.items()
        if isinstance(values, dict) and name not in ("macro avg", "weighted avg")
    ]
    weekly_dirs = sorted((CATEGORY_TUNING_DIR / "tuning_outputs").glob("weekly_*"))
    decision = _read_json(weekly_dirs[-1] / "model_promotion_decision.json") if weekly_dirs else {}
    label_files = [
        path
        for folder in ("CFE_input", "CFE_reviewed_issue")
        for path in (CATEGORY_TUNING_DIR / folder).glob("*.csv")
        if path.resolve() != CATEGORY_REVIEW_PATH.resolve() and path.name.startswith(("weekly_human_labels", "reviewed"))
    ]
    latest_labels = max(label_files, key=lambda path: path.stat().st_mtime, default=None)
    return {
        "available": bool(metrics),
        "model_file": CATEGORY_MODEL_PATH.name,
        "trained_at": _file_date(CATEGORY_METRICS_PATH) or _file_date(CATEGORY_MODEL_PATH),
        "accuracy": metrics.get("accuracy"),
        "macro_f1": (report.get("macro avg") or {}).get("f1-score"),
        "rows_used": metrics.get("rows_used"),
        "test_rows": metrics.get("test_rows_human") or (report.get("macro avg") or {}).get("support"),
        "dropped_rare_labels": metrics.get("dropped_rare_labels") or [],
        "per_category": per_category,
        "latest_weekly_run": weekly_dirs[-1].name.removeprefix("weekly_") if weekly_dirs else "",
        "latest_human_responses": decision.get("human_responses_received"),
        "latest_promoted": decision.get("promoted"),
        "latest_human_labels_at": _file_date(latest_labels) if latest_labels else "",
    }


def _read_category_reviews() -> dict[str, dict]:
    if not CATEGORY_REVIEW_PATH.is_file():
        return {}
    with CATEGORY_REVIEW_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        return {row["ips_case_number"]: row for row in csv.DictReader(handle) if row.get("ips_case_number")}


def _write_category_reviews(reviews: dict[str, dict]) -> None:
    CATEGORY_REVIEW_PATH.parent.mkdir(parents=True, exist_ok=True)
    temp_path = CATEGORY_REVIEW_PATH.with_suffix(".tmp")
    with temp_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CATEGORY_REVIEW_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(sorted(reviews.values(), key=lambda row: row.get("reviewed_at", "")))
    os.replace(temp_path, CATEGORY_REVIEW_PATH)


def _category_technology_expr(available: set[str]) -> str:
    if "technology" in available:
        return "technology::text"
    if "bug_project" not in available:
        return "NULL"
    # Same bug_project -> technology mapping the offload weighting uses.
    return (
        "CASE LOWER(TRIM(COALESCE(bug_project::text, ''))) "
        "WHEN 'wifi' THEN 'WiFi' WHEN 'bt' THEN 'BT' WHEN 'cie' THEN 'Software' WHEN 'wot' THEN 'Tools' ELSE NULL END"
    )


def _load_category_queue() -> dict[str, dict]:
    model_mtime = CATEGORY_MODEL_PATH.stat().st_mtime
    if CATEGORY_CACHE["model_mtime"] != model_mtime:
        CATEGORY_CACHE.update(
            model_mtime=model_mtime,
            bundle=category_model.load_category_model(str(CATEGORY_MODEL_PATH)),
            predictions={},
        )

    db = Sherlock.PostgresCustomerEngineeringDb
    connection = psycopg2.connect(
        dbname=db.database, user=db.user, password=db.password, host=db.host, port=db.port, connect_timeout=10
    )
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT column_name FROM information_schema.columns "
                "WHERE table_name = %s AND table_schema = ANY(current_schemas(false))",
                (FIELD_ISSUE_TABLE,),
            )
            available = {row[0] for row in cursor.fetchall()}
            cursor.execute(
                sql.SQL(
                    "SELECT DISTINCT ON (ips_case_number::text) ips_case_number::text, ips_title, ips_url, ips_status, "
                    "engineer, ips_created_date, {technology} FROM {table} "
                    "WHERE ips_created_date >= make_date(%s, 1, 1) AND ips_created_date < make_date(%s, 1, 1) "
                    "AND ips_case_number::text ~ '^[0-9]+$' "
                    "AND NULLIF(NULLIF(TRIM(ips_title::text), ''), 'NA') IS NOT NULL "
                    "ORDER BY ips_case_number::text, ips_created_date DESC"
                ).format(
                    technology=sql.SQL(_category_technology_expr(available)),
                    table=sql.Identifier(FIELD_ISSUE_TABLE),
                ),
                (CURRENT_YEAR, CURRENT_YEAR + 1),
            )
            rows = cursor.fetchall()
            cursor.execute(
                sql.SQL(
                    "SELECT DISTINCT TRIM(engineer::text) FROM {table} "
                    "WHERE NULLIF(NULLIF(TRIM(engineer::text), ''), 'NA') IS NOT NULL"
                ).format(table=sql.Identifier(FIELD_ISSUE_TABLE))
            )
            reviewers = {row[0] for row in cursor.fetchall()}
    finally:
        connection.close()

    queue: dict[str, dict] = {}
    for case_number, title, url, status, engineer, created, technology in rows:
        technology = _field_issue_text(technology)
        # Software issues are hard-ruled to ICPS/Killer in both prediction and training, so reviews would be ignored.
        if technology.lower() == "software":
            continue
        title = _field_issue_text(title)
        key = (case_number, title, technology)
        if key not in CATEGORY_CACHE["predictions"]:
            CATEGORY_CACHE["predictions"][key] = category_model.classify_issue_title(
                CATEGORY_CACHE["bundle"], title, technology=technology, use_llm=False
            )
        predicted, confidence = CATEGORY_CACHE["predictions"][key]
        queue[case_number] = {
            "ips_case_number": case_number,
            "ips_title": title,
            "ips_url": _field_issue_text(url),
            "ips_status": _field_issue_text(status),
            "engineer": _field_issue_text(engineer) or "Unassigned",
            "created": created.date().isoformat() if created else "",
            "technology": technology,
            "predicted_category": predicted,
            "confidence": round(float(confidence), 4),
        }
    CATEGORY_CACHE["queue"] = queue
    CATEGORY_CACHE["reviewers"] = reviewers
    CATEGORY_QUEUE_STATE["loaded_at"] = datetime.now().timestamp()
    return queue


def _refresh_category_queue() -> None:
    try:
        _load_category_queue()
    except Exception as exc:
        print(f"[dashboard] Category queue refresh failed ({type(exc).__name__}); keeping previous data.")
    finally:
        CATEGORY_QUEUE_STATE["refreshing"] = False


def _cached_category_queue() -> dict[str, dict]:
    if not CATEGORY_QUEUE_STATE["loaded_at"]:
        return _load_category_queue()
    if datetime.now().timestamp() - CATEGORY_QUEUE_STATE["loaded_at"] > CATEGORY_QUEUE_MAX_AGE:
        with CATEGORY_QUEUE_LOCK:
            if not CATEGORY_QUEUE_STATE["refreshing"]:
                CATEGORY_QUEUE_STATE["refreshing"] = True
                threading.Thread(target=_refresh_category_queue, daemon=True).start()
    return CATEGORY_CACHE["queue"]


def _llm_predictions() -> dict[str, dict]:
    with LLM_LOCK:
        if LLM_STATE["predictions"] is None:
            LLM_STATE["predictions"] = _read_json(LLM_CACHE_PATH)
        return LLM_STATE["predictions"]


def _save_llm_predictions() -> None:
    with LLM_LOCK:
        snapshot = json.dumps(LLM_STATE["predictions"] or {}, ensure_ascii=False)
    LLM_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    temp_path = LLM_CACHE_PATH.with_suffix(".tmp")
    temp_path.write_text(snapshot, encoding="utf-8")
    os.replace(temp_path, LLM_CACHE_PATH)


def _llm_prediction_for(issue: dict) -> dict | None:
    cached = _llm_predictions().get(issue["ips_case_number"])
    # A title/technology edit or a different LLM model invalidates the cached answer.
    if not cached or cached.get("title") != issue["ips_title"] or cached.get("technology") != issue["technology"]:
        return None
    return cached if cached.get("llm_model") == category_model.llm_model_name() else None


def _run_llm_backfill(missing: list[dict], llm_model: str) -> None:
    bundle = CATEGORY_CACHE["bundle"]

    def predict(issue: dict) -> tuple[dict, tuple[str, float] | None]:
        return issue, category_model.classify_issue_title_llm(bundle, issue["ips_title"], technology=issue["technology"])

    def record(issue: dict, result: tuple[str, float] | None) -> None:
        with LLM_LOCK:
            LLM_STATE["done"] += 1
            if result is None:
                LLM_STATE["failed"] += 1
                return
            LLM_STATE["predictions"][issue["ips_case_number"]] = {
                "title": issue["ips_title"],
                "technology": issue["technology"],
                "category": result[0],
                "confidence": round(float(result[1]), 4),
                "llm_model": llm_model,
                "predicted_at": datetime.now().isoformat(timespec="seconds"),
            }

    error = ""
    try:
        consecutive_failures = 0
        batch_size = 4
        with ThreadPoolExecutor(max_workers=batch_size) as pool:
            for start in range(0, len(missing), batch_size):
                for issue, result in pool.map(predict, missing[start : start + batch_size]):
                    record(issue, result)
                    consecutive_failures = consecutive_failures + 1 if result is None else 0
                if consecutive_failures >= LLM_MAX_CONSECUTIVE_FAILURES:
                    error = "LLM requests keep failing (for example HTTP 401/403). Check that GNAI_TOKEN in .env is valid and not expired, then restart the server."
                    break
                # Persist every batch so a crash or restart never re-spends tokens on finished issues.
                _save_llm_predictions()
    finally:
        _save_llm_predictions()
        with LLM_LOCK:
            LLM_STATE.update(running=False, error=error, failed_at=datetime.now() if error else None)


def _start_llm_backfill(queue: dict[str, dict]) -> None:
    llm_model = category_model.llm_model_name()
    if not llm_model:
        return
    missing = [issue for issue in queue.values() if _llm_prediction_for(issue) is None]
    with LLM_LOCK:
        if LLM_STATE["running"] or not missing:
            return
        if LLM_STATE["failed_at"] and datetime.now() - LLM_STATE["failed_at"] < LLM_RETRY_AFTER:
            return
        LLM_STATE.update(running=True, done=0, total=len(missing), failed=0, error="")
    threading.Thread(target=_run_llm_backfill, args=(missing, llm_model), daemon=True).start()


def _llm_status() -> dict:
    with LLM_LOCK:
        return {
            "model": category_model.llm_model_name(),
            "running": LLM_STATE["running"],
            "done": LLM_STATE["done"],
            "total": LLM_STATE["total"],
            "failed": LLM_STATE["failed"],
            "error": LLM_STATE["error"],
        }


def _llm_payload() -> dict:
    predictions = {}
    for case_number, issue in CATEGORY_CACHE["queue"].items():
        cached = _llm_prediction_for(issue)
        if cached:
            predictions[case_number] = {"category": cached["category"], "confidence": cached["confidence"]}
    return {"llm": _llm_status(), "predictions": predictions}


def _load_category_prediction() -> dict:
    queue = _cached_category_queue()
    _start_llm_backfill(queue)
    reviews = _read_category_reviews()
    issues = []
    for case_number, issue in queue.items():
        review = reviews.get(case_number)
        llm = _llm_prediction_for(issue)
        issues.append(
            {
                **issue,
                "llm_category": llm["category"] if llm else None,
                "llm_confidence": llm["confidence"] if llm else None,
                "review": {
                    key: review.get(key, "")
                    for key in ("human_category", "verdict", "llm_category", "reviewer", "reviewed_at")
                }
                if review
                else None,
            }
        )
    return {
        "year": CURRENT_YEAR,
        "loaded_at": datetime.now().isoformat(timespec="seconds"),
        "model": _category_model_status(),
        "llm": _llm_status(),
        "categories": _category_labels(),
        "reviewers": sorted(CATEGORY_CACHE["reviewers"], key=str.casefold),
        "review_file": CATEGORY_REVIEW_PATH.name,
        "issues": issues,
    }


def _save_category_review(payload: dict) -> dict:
    case_number = str(payload.get("ips_case_number") or "").strip()
    human_category = str(payload.get("human_category") or "").strip()
    reviewer = str(payload.get("reviewer") or "").strip()
    if case_number not in CATEGORY_CACHE["queue"]:
        _load_category_queue()
    issue = CATEGORY_CACHE["queue"].get(case_number)
    if issue is None:
        raise ValueError("This issue is not in the current review list. Reload the page.")
    if reviewer not in CATEGORY_CACHE["reviewers"]:
        raise ValueError("Choose your name under 'Reviewing as' before saving.")
    if human_category and human_category not in _category_labels():
        raise ValueError("Pick a category from the list.")

    with CATEGORY_LOCK:
        reviews = _read_category_reviews()
        if not human_category:
            reviews.pop(case_number, None)
            _write_category_reviews(reviews)
            return {"ips_case_number": case_number, "review": None}
        llm = _llm_prediction_for(issue)
        review = {
            "ips_title": issue["ips_title"],
            "predicted_category": issue["predicted_category"],
            "technology": issue["technology"],
            "human_category": human_category,
            "ips_case_number": case_number,
            "verdict": "correct" if human_category == issue["predicted_category"] else "corrected",
            "llm_category": llm["category"] if llm else "",
            "reviewer": reviewer,
            "reviewed_at": datetime.now().isoformat(timespec="seconds"),
        }
        reviews[case_number] = review
        _write_category_reviews(reviews)
    return {
        "ips_case_number": case_number,
        "review": {key: review[key] for key in ("human_category", "verdict", "llm_category", "reviewer", "reviewed_at")},
    }


class LoadingDashboardHandler(SimpleHTTPRequestHandler):
    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        requested_path = Path(self.translate_path(self.path)).resolve()
        # Never serve dotfiles such as .env (DB credentials) or .venv.
        if requested_path in BLOCKED_STATIC_PATHS or any(
            part.startswith(".") for part in unquote(parsed.path).split("/")
        ):
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        if parsed.path == "/api/activity":
            self._send_activity(parse_qs(parsed.query).get("refresh", ["0"])[0] == "1")
            return
        if parsed.path == "/api/field-issues":
            self._send_field_issues()
            return
        if parsed.path == "/api/category-prediction":
            self._send_category_prediction()
            return
        if parsed.path == "/api/category-llm":
            self._send_json(_llm_payload())
            return
        if parsed.path == "/api/loading":
            self._send_snapshot(parse_qs(parsed.query).get("date", ["latest"])[0])
            return
        if parsed.path == "/api/dates":
            self._send_dates()
            return
        if parsed.path == "/api/weights":
            self._send_weights()
            return
        if parsed.path == "/api/batch":
            self._send_batch_status()
            return
        if parsed.path in ("", "/"):
            self.path = "/index.html"
        super().do_GET()

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/run-batch":
            self._start_batch()
            return
        if parsed.path == "/api/restart":
            self._restart_server()
            return
        if parsed.path == "/api/category-review":
            self._save_category_review()
            return
        self.send_error(HTTPStatus.NOT_FOUND)

    def do_OPTIONS(self) -> None:
        self.send_response(HTTPStatus.NO_CONTENT)
        self._add_cors_headers()
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type, Authorization")
        self.send_header("Content-Length", "0")
        self.end_headers()

    def _add_cors_headers(self) -> None:
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type, Authorization")

    def _send_json(self, payload: dict, status: HTTPStatus = HTTPStatus.OK) -> None:
        body = json.dumps(payload, ensure_ascii=True).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self._add_cors_headers()
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _read_snapshot(self, selected_date: str) -> dict | None:
        path = SNAPSHOT_PATH if selected_date in ("", "latest") else HISTORY_PATH / f"{selected_date}.json"
        try:
            with path.open("r", encoding="utf-8") as handle:
                snapshot = json.load(handle)
            return snapshot if isinstance(snapshot, dict) else None
        except (OSError, ValueError):
            if selected_date not in ("", "latest"):
                latest = self._read_snapshot("latest")
                if latest and str(latest.get("generated_at", ""))[:10] == selected_date:
                    return latest
            return None

    def _send_dates(self) -> None:
        dates = []
        if HISTORY_PATH.exists():
            dates = sorted(
                (path.stem for path in HISTORY_PATH.glob("????-??-??.json")),
                reverse=True,
            )
        latest = self._read_snapshot("latest")
        latest_date = str(latest.get("generated_at", ""))[:10] if latest else ""
        if latest_date and latest_date not in dates:
            dates.insert(0, latest_date)
        self._send_json({"dates": dates})

    def _send_weights(self) -> None:
        payload = {"default_weight": 1.0, "category_weights": {}, "category_technology_weights": {}}
        try:
            with WEIGHT_MAP_PATH.open("r", encoding="utf-8") as handle:
                source = json.load(handle)
            if isinstance(source, dict):
                payload.update(
                    {
                        "default_weight": source.get("default_weight", 1.0),
                        "category_weights": source.get("category_weights") or {},
                        "category_technology_weights": source.get("category_technology_weights") or {},
                    }
                )
        except (OSError, ValueError):
            payload["message"] = "Weight configuration is unavailable."
        self._send_json(payload)

    def _send_snapshot(self, selected_date: str) -> None:
        payload = {
            "generated_at": None,
            "categories": [],
            "rows": [],
            "available": False,
            "selected_date": selected_date,
            "message": "No batch result is available yet. Run the daily loading batch first.",
        }
        snapshot = self._read_snapshot(selected_date)
        if snapshot is not None:
            payload.update(
                {
                    "generated_at": snapshot.get("generated_at"),
                    "categories": snapshot.get("categories") or [],
                    "rows": snapshot.get("rows") or [],
                    "available": True,
                    "message": "",
                }
            )

        self._send_json(payload)

    def _send_activity(self, force: bool = False) -> None:
        try:
            self._send_json(_load_activity_workbook(force))
        except FileNotFoundError as exc:
            self._send_json(
                {"available": False, "message": f"Workbook unavailable: {exc}"},
                HTTPStatus.OK,
            )
        except ValueError as exc:
            self._send_json({"available": False, "message": f"Workbook unavailable: {exc}"}, HTTPStatus.OK)
        except Exception as exc:
            self._send_json(
                {"available": False, "message": f"Unable to read workbook ({type(exc).__name__})."},
                HTTPStatus.INTERNAL_SERVER_ERROR,
            )

    def _send_field_issues(self) -> None:
        try:
            self._send_json(_load_field_issues())
        except Exception as exc:
            self._send_json(
                {"issues": [], "message": f"Unable to query {FIELD_ISSUE_TABLE} ({type(exc).__name__})."},
                HTTPStatus.INTERNAL_SERVER_ERROR,
            )

    def _send_category_prediction(self) -> None:
        try:
            self._send_json(_load_category_prediction())
        except FileNotFoundError:
            self._send_json(
                {"issues": [], "message": f"Category model not found: {CATEGORY_MODEL_PATH.name}."},
                HTTPStatus.NOT_FOUND,
            )
        except Exception as exc:
            self._send_json(
                {"issues": [], "message": f"Unable to load predictions ({type(exc).__name__})."},
                HTTPStatus.INTERNAL_SERVER_ERROR,
            )

    def _save_category_review(self) -> None:
        length = int(self.headers.get("Content-Length") or 0)
        if not 0 < length <= 4096:
            self._send_json({"message": "Invalid request body."}, HTTPStatus.BAD_REQUEST)
            return
        try:
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            if not isinstance(payload, dict):
                raise ValueError("Invalid request body.")
            self._send_json(_save_category_review(payload))
        except ValueError as exc:
            self._send_json({"message": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:
            self._send_json({"message": f"Unable to save review ({type(exc).__name__})."}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def _send_batch_status(self) -> None:
        global RUN_FINISHED_AT, RUN_RETURN_CODE
        with RUN_LOCK:
            process = RUN_PROCESS
            if process is not None and process.poll() is not None:
                RUN_RETURN_CODE = process.returncode
                if RUN_FINISHED_AT is None:
                    RUN_FINISHED_AT = datetime.now().isoformat(timespec="seconds")
            running = process is not None and process.poll() is None
            payload = {
                "running": running,
                "started_at": RUN_STARTED_AT,
                "finished_at": RUN_FINISHED_AT,
                "return_code": RUN_RETURN_CODE,
            }
        self._send_json(payload)

    def _start_batch(self) -> None:
        global RUN_PROCESS, RUN_STARTED_AT, RUN_FINISHED_AT, RUN_RETURN_CODE
        with RUN_LOCK:
            if RUN_PROCESS is not None and RUN_PROCESS.poll() is None:
                self._send_json({"running": True, "message": "The batch is already running."}, HTTPStatus.CONFLICT)
                return
            if not BATCH_PATH.exists():
                self._send_json({"running": False, "message": f"Batch file not found: {BATCH_PATH.name}"}, HTTPStatus.NOT_FOUND)
                return
            if os.name == "nt":
                command = ["cmd.exe", "/d", "/c", f"{BATCH_PATH} --no-email"]
            else:
                command = ["sh", str(BATCH_PATH), "--no-email"]
            RUN_PROCESS = subprocess.Popen(command, cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            RUN_STARTED_AT = datetime.now().isoformat(timespec="seconds")
            RUN_FINISHED_AT = None
            RUN_RETURN_CODE = None
        self._send_json({"running": True, "started_at": RUN_STARTED_AT}, HTTPStatus.ACCEPTED)

    def _restart_server(self) -> None:
        global RESTART_REQUESTED_AT
        if os.name != "nt" or not RESTART_SCRIPT_PATH.is_file():
            self._send_json({"message": "Restart is only available on the Windows server."}, HTTPStatus.BAD_REQUEST)
            return
        with RUN_LOCK:
            now = datetime.now()
            if RESTART_REQUESTED_AT and now - RESTART_REQUESTED_AT < RESTART_COOLDOWN:
                self._send_json({"message": "A restart was requested less than 2 minutes ago."}, HTTPStatus.CONFLICT)
                return
            # Launch via WMI so the script is not a child of this process and survives the task stop it performs.
            inner = f'powershell.exe -NoProfile -ExecutionPolicy Bypass -WindowStyle Hidden -File "{RESTART_SCRIPT_PATH}"'
            launcher = (
                "$r = Invoke-CimMethod -ClassName Win32_Process -MethodName Create "
                f"-Arguments @{{CommandLine='{inner.replace(chr(39), chr(39) * 2)}'}}; exit $r.ReturnValue"
            )
            try:
                result = subprocess.run(
                    ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", launcher],
                    capture_output=True,
                    timeout=30,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired) as exc:
                self._send_json({"message": f"Unable to start restart ({type(exc).__name__})."}, HTTPStatus.INTERNAL_SERVER_ERROR)
                return
            if result.returncode != 0:
                self._send_json({"message": f"Unable to start restart (code {result.returncode})."}, HTTPStatus.INTERNAL_SERVER_ERROR)
                return
            RESTART_REQUESTED_AT = now
        self._send_json({"restarting": True}, HTTPStatus.ACCEPTED)

    def log_message(self, format: str, *args: object) -> None:
        print(f"[dashboard] {self.address_string()} - {format % args}")


def _warm_caches() -> None:
    """Pre-load the slow data at startup so the first visitor does not wait for it."""
    for name, loader in (("activity workbook", _load_activity_workbook), ("category queue", _cached_category_queue)):
        try:
            loader()
            print(f"[dashboard] Warmed {name} cache.")
        except Exception as exc:
            print(f"[dashboard] Could not warm {name} cache ({type(exc).__name__}).")


def main() -> None:
    parser = argparse.ArgumentParser(description="Serve the weighted team loading dashboard.")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8010)
    args = parser.parse_args()
    server = ThreadingHTTPServer((args.host, args.port), LoadingDashboardHandler)
    print(f"Loading dashboard: http://{args.host}:{args.port}/")
    threading.Thread(target=_warm_caches, daemon=True).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
